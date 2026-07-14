"""Excel worksheet table sync for AREV2 writer output."""
from __future__ import annotations

import re
from typing import List

from openpyxl import Workbook
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.worksheet.worksheet import Worksheet

_TABLE_STYLE = TableStyleInfo(
    name="TableStyleMedium2",
    showFirstColumn=False,
    showLastColumn=False,
    showRowStripes=True,
    showColumnStripes=False,
)
_ARE_PREFIX = "ARE_"


def headers_table_eligible(headers: List[str]) -> bool:
    if not headers:
        return False
    normalized = [str(h).strip() if h is not None else "" for h in headers]
    if any(not value for value in normalized):
        return False
    return len(normalized) == len(set(normalized))


def sync_worksheet_table(ws: Worksheet, wb: Workbook, headers: List[str], sheet_name: str) -> None:
    """Create or extend an Excel table when headers and data rows are valid."""
    if not headers_table_eligible(headers):
        return
    if ws.max_row < 2:
        return

    last_col = get_column_letter(len(headers))
    ref = f"A1:{last_col}{ws.max_row}"

    existing = _find_managed_table(ws)
    if existing is not None:
        existing.ref = ref
        return

    display_name = _unique_workbook_table_name(wb, _table_base_name(sheet_name))
    table = Table(displayName=display_name, ref=ref)
    table.tableStyleInfo = _TABLE_STYLE
    ws.add_table(table)


def _table_base_name(sheet_name: str) -> str:
    cleaned = re.sub(r"[^A-Za-z0-9_]", "_", str(sheet_name or "")).strip("_")
    if not cleaned:
        return f"{_ARE_PREFIX}Table"
    name = f"{_ARE_PREFIX}{cleaned}"
    if not (name[0].isalpha() or name[0] == "_"):
        name = f"{_ARE_PREFIX}Table"
    return name[:240]


def _workbook_table_names(wb: Workbook) -> set[str]:
    names: set[str] = set()
    for worksheet in wb.worksheets:
        for table in worksheet.tables.values():
            names.add(str(table.displayName))
    return names


def _unique_workbook_table_name(wb: Workbook, base: str) -> str:
    used = _workbook_table_names(wb)
    candidate = base
    suffix = 2
    while candidate in used:
        candidate = f"{base}_{suffix}"
        suffix += 1
    return candidate


def _find_managed_table(ws: Worksheet) -> Table | None:
    are_tables = [table for table in ws.tables.values() if str(table.displayName).startswith(_ARE_PREFIX)]
    if len(are_tables) == 1:
        return are_tables[0]
    if len(are_tables) > 1:
        return are_tables[0]

    a1_tables: list[Table] = []
    for table in ws.tables.values():
        ref = str(table.ref or "").upper()
        if ref.startswith("A1:"):
            a1_tables.append(table)
    if len(a1_tables) == 1:
        return a1_tables[0]
    return None
