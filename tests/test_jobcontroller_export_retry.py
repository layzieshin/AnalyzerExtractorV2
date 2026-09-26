import json
import sqlite3
from pathlib import Path

import pytest
from openpyxl import load_workbook

from src.jobcontroller.api import release_validated_run, retry_excel_export, submit
from src.jobcontroller.jobcontroller import JobController
from src.parser.model import ParsedDocument, ParsedPage


def _ruleset(assay_key: str, assay_name: str, filename: str) -> dict:
    return {
        "assay_name": assay_name,
        "assay_key": assay_key,
        "lot_rule": {"regex": r"Lot:\s*(\w+)"},
        "extract_rules": {
            "fields": [
                {"key": "plate_name", "regex": r"Plate:\s*(\w+)", "required": True},
                {"key": "test", "regex": r"Test:\s*(\w+)", "required": True},
                {"key": "date", "regex": r"Date:\s*([0-9\-]+)", "required": True},
                {"key": "time", "regex": r"Time:\s*([0-9:]+)", "required": True},
            ]
        },
        "excel_rules": {
            "excel_filename_template": filename,
            "sheetname_template": "{lot_id}",
            "column_mapping": {"test": "Test"},
        },
    }


def _write_project(root: Path, *, assays: list[tuple[str, str, str, str]]) -> None:
    (root / "rules").mkdir(parents=True)
    (root / "output" / "final").mkdir(parents=True)
    index = {
        "assays": [
            {"assay_key": assay_key, "ruleset_file": ruleset_file}
            for assay_key, _name, ruleset_file, _filename in assays
        ]
    }
    (root / "rules" / "index.json").write_text(json.dumps(index), encoding="utf-8")
    for assay_key, assay_name, ruleset_file, filename in assays:
        (root / "rules" / ruleset_file).write_text(
            json.dumps(_ruleset(assay_key, assay_name, filename)),
            encoding="utf-8",
        )


def _lines(assay_name: str, assay_key: str, lot: str, plate: str, test: str) -> list[str]:
    return [
        assay_name,
        assay_key,
        f"Lot: {lot}",
        f"Plate: {plate}",
        f"Test: {test}",
        "Date: 2026-05-28",
        "Time: 12:00:00",
    ]


def _parse_one(pdf: Path) -> ParsedDocument:
    lines = _lines("Assay A", "(1111)", "LOTA", "PLATE1", "TESTX")
    return ParsedDocument(source_path=str(pdf), pages=[ParsedPage(page_number=1, lines=lines)], meta={"page_count": 1})


def _parse_two(pdf: Path) -> ParsedDocument:
    lines = _lines("Assay A", "(1111)", "LOTA", "PLATE1", "TESTX") + _lines(
        "Assay B", "(2222)", "LOTB", "PLATE2", "TESTY"
    )
    return ParsedDocument(source_path=str(pdf), pages=[ParsedPage(page_number=1, lines=lines)], meta={"page_count": 1})


def _state(root: Path, job_id: str) -> dict:
    return json.loads((root / "jobs" / f"{job_id}.json").read_text(encoding="utf-8"))


def _run_count(sqlite_path: Path) -> int:
    with sqlite3.connect(sqlite_path) as conn:
        return int(conn.execute("SELECT COUNT(*) FROM runs").fetchone()[0])


def _run_ids(sqlite_path: Path) -> list[int]:
    with sqlite3.connect(sqlite_path) as conn:
        return [int(row[0]) for row in conn.execute("SELECT id FROM runs ORDER BY id").fetchall()]


def test_both_excel_failure_persists_plan_after_single_sqlite_run(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    root = tmp_path / "proj"
    _write_project(root, assays=[("(1111)", "Assay A", "AssayA.json", "{assay_name}.xlsx")])
    pdf = root / "input.pdf"
    pdf.write_bytes(b"both-fail")
    sqlite_path = root / "output" / "final" / "res.sqlite3"
    calls = {"sqlite": 0, "excel": 0}

    monkeypatch.setattr("src.parser.api.parse", lambda path: _parse_one(Path(path)))

    real_sqlite = __import__("src.dbwriter.api", fromlist=["write_record_sqlite"]).write_record_sqlite
    real_excel = __import__("src.writer.api", fromlist=["write_record"]).write_record

    def sqlite_once(*args, **kwargs):
        calls["sqlite"] += 1
        return real_sqlite(*args, **kwargs)

    def excel_fail(*_args, **_kwargs):
        calls["excel"] += 1
        raise RuntimeError("disk full")

    monkeypatch.setattr("src.dbwriter.api.write_record_sqlite", sqlite_once)
    monkeypatch.setattr("src.writer.api.write_record", excel_fail)

    result = submit(str(pdf), str(root), output_mode="both", sqlite_path=str(sqlite_path), device_id="dev1")

    assert result.status == "DONE"
    assert calls == {"sqlite": 1, "excel": 0}
    assert _run_count(sqlite_path) == 1
    state = _state(root, result.job_id)
    assert state["status"] == "DONE"
    assert state["sqlite_status"] == "success"
    assert state["structured_status"] == "success"
    assert state["excel_status"] == "awaiting_validation"
    assert state["export_retry_available"] is False
    plan = state["output_plan"]
    assert plan["version"] == 2
    assert plan["output_mode"] == "both"
    assay = plan["assays"][0]
    assert assay["excel_status"] == "awaiting_validation"
    assert assay["sqlite_run_id"] == _run_ids(sqlite_path)[0]
    assert assay["assay_record"]["dedupe_key"]
    assert assay["ruleset"]["excel_rules"]["excel_filename_template"] == "{assay_name}.xlsx"
    assert assay["write_item"]["outputs"][0]["sink"] == "sqlite"
    assert not list((root / "jobs").glob("*.tmp"))
    assert not list((root / "jobs").glob(".*.tmp"))


def test_export_retry_calls_writer_only_and_keeps_plan_when_it_fails(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    root = tmp_path / "proj"
    _write_project(root, assays=[("(1111)", "Assay A", "AssayA.json", "{assay_name}.xlsx")])
    pdf = root / "input.pdf"
    pdf.write_bytes(b"retry-fail")
    sqlite_path = root / "output" / "final" / "res.sqlite3"
    pipeline = {
        "parse": 0,
        "normalize": 0,
        "choose": 0,
        "split": 0,
        "extract": 0,
        "sqlite": 0,
        "rules": 0,
        "revalidation": 0,
        "excel": 0,
    }

    monkeypatch.setattr("src.parser.api.parse", lambda path: _parse_one(Path(path)))
    real_sqlite = __import__("src.dbwriter.api", fromlist=["write_record_sqlite"]).write_record_sqlite

    def counting_sqlite(*args, **kwargs):
        pipeline["sqlite"] += 1
        return real_sqlite(*args, **kwargs)

    monkeypatch.setattr("src.dbwriter.api.write_record_sqlite", counting_sqlite)
    monkeypatch.setattr(
        "src.writer.api.write_record",
        lambda *_a, **_k: (_ for _ in ()).throw(RuntimeError("locked")),
    )
    first = submit(str(pdf), str(root), output_mode="both", sqlite_path=str(sqlite_path), device_id="dev1")
    assert first.status == "DONE"
    ids = _run_ids(sqlite_path)
    released = release_validated_run(
        root, first.job_id, ids[0], operator_initials="AB",
        validated_at_utc="2026-09-25T10:00:00+00:00", validation_id=1,
    )
    assert released.status == "FAILED"
    dumps = {
        path.name: path.read_bytes()
        for path in (root / "jobs").glob(f"{first.job_id}*")
        if path.suffix != ".json"
    }

    def forbid(name):
        def _forbidden(*_args, **_kwargs):
            pipeline[name] += 1
            raise AssertionError(name)

        return _forbidden

    monkeypatch.setattr("src.parser.api.parse", forbid("parse"))
    monkeypatch.setattr("src.normalizer.api.normalize_lines", forbid("normalize"))
    monkeypatch.setattr("src.assaychooser.api.detect_assays", forbid("choose"))
    monkeypatch.setattr("src.contentsplitter.api.split_by_assay_name_and_key", forbid("split"))
    monkeypatch.setattr("src.extractor.api.extract_record", forbid("extract"))
    monkeypatch.setattr("src.dbwriter.api.write_record_sqlite", forbid("sqlite"))
    monkeypatch.setattr("src.ruleresolver.api.resolve_ruleset", forbid("rules"))
    monkeypatch.setattr("src.resultvalidation.api.get_run_validation", forbid("revalidation"))
    monkeypatch.setattr(
        "src.writer.api.write_record",
        lambda *_a, **_k: (_ for _ in ()).throw(RuntimeError("still locked")),
    )

    pdf.unlink()
    second = retry_excel_export(root, first.job_id)
    assert second.status == "FAILED"
    assert pipeline["parse"] == 0
    assert pipeline["normalize"] == 0
    assert pipeline["choose"] == 0
    assert pipeline["split"] == 0
    assert pipeline["extract"] == 0
    assert pipeline["sqlite"] == 1
    assert pipeline["rules"] == 0
    assert pipeline["revalidation"] == 1
    state = _state(root, first.job_id)
    assert state["export_retry_available"] is True
    assert state["sqlite_status"] == "success"
    assert state["structured_status"] == "success"
    assert state["excel_status"] == "failed"
    assert state["output_plan"]["output_mode"] == "both"
    assert state["output_plan"]["sqlite_status"] == "success"
    assert state["output_plan"]["version"] == 2
    assert _run_ids(sqlite_path) == ids
    assert dumps == {
        path.name: path.read_bytes()
        for path in (root / "jobs").glob(f"{first.job_id}*")
        if path.suffix != ".json"
    }


def test_successful_export_retry_does_not_write_sqlite_again(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    root = tmp_path / "proj"
    _write_project(root, assays=[("(1111)", "Assay A", "AssayA.json", "Frozen.xlsx")])
    pdf = root / "input.pdf"
    pdf.write_bytes(b"retry-ok")
    sqlite_path = root / "output" / "final" / "res.sqlite3"
    excel_calls = {"count": 0}
    monkeypatch.setattr("src.parser.api.parse", lambda path: _parse_one(Path(path)))
    real_excel = __import__("src.writer.api", fromlist=["write_record"]).write_record

    def fail_once(record, ruleset, output_dir, **kwargs):
        excel_calls["count"] += 1
        if excel_calls["count"] == 1:
            raise RuntimeError("busy")
        return real_excel(record, ruleset, output_dir, **kwargs)

    monkeypatch.setattr("src.writer.api.write_record", fail_once)
    first = submit(str(pdf), str(root), output_mode="both", sqlite_path=str(sqlite_path), device_id="dev1")
    assert first.status == "DONE"
    run_id = _run_ids(sqlite_path)[0]
    released = release_validated_run(
        root, first.job_id, run_id, operator_initials="AB",
        validated_at_utc="2026-09-25T10:00:00+00:00", validation_id=1,
    )
    assert released.status == "FAILED"
    (root / "rules" / "AssayA.json").write_text(
        json.dumps(_ruleset("(1111)", "Assay A", "Changed.xlsx")),
        encoding="utf-8",
    )
    monkeypatch.setattr("src.parser.api.parse", lambda *_a, **_k: (_ for _ in ()).throw(AssertionError("parse")))
    monkeypatch.setattr(
        "src.dbwriter.api.write_record_sqlite",
        lambda *_a, **_k: (_ for _ in ()).throw(AssertionError("sqlite")),
    )
    pdf.unlink()
    second = retry_excel_export(root, first.job_id)
    assert second.status == "DONE"
    assert _run_count(sqlite_path) == 1
    assert (root / "output" / "final" / "Frozen.xlsx").is_file()
    assert not (root / "output" / "final" / "Changed.xlsx").exists()
    state = _state(root, first.job_id)
    assert state["status"] == "DONE"
    assert state["export_retry_available"] is False
    assert state["excel_status"] == "success"
    assert state["sqlite_status"] == "success"
    assert state["structured_status"] == "success"
    assert state["output_plan"]["output_mode"] == "both"
    corrected = release_validated_run(
        root, first.job_id, run_id, operator_initials="CD",
        validated_at_utc="2026-09-25T11:00:00+00:00", validation_id=2,
        correction=True,
    )
    assert corrected.status == "DONE"
    wb = load_workbook(root / "output" / "final" / "Frozen.xlsx")
    ws = wb["LOTA"]
    headers = [cell.value for cell in ws[1]]
    assert ws.max_row == 2
    assert ws.cell(2, headers.index("VALIDIERT_DURCH") + 1).value == "CD"
    assert ws.cell(2, headers.index("VALIDIERT_AM") + 1).value == "2026-09-25T11:00:00+00:00"
    assert not pdf.exists()


def test_multi_assay_sqlite_precedes_excel_and_retry_writes_only_missing(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    root = tmp_path / "proj"
    _write_project(
        root,
        assays=[
            ("(1111)", "Assay A", "AssayA.json", "A.xlsx"),
            ("(2222)", "Assay B", "AssayB.json", "B.xlsx"),
        ],
    )
    pdf = root / "input.pdf"
    pdf.write_bytes(b"multi")
    sqlite_path = root / "output" / "final" / "res.sqlite3"
    order: list[str] = []
    monkeypatch.setattr("src.parser.api.parse", lambda path: _parse_two(Path(path)))
    real_sqlite = __import__("src.dbwriter.api", fromlist=["write_record_sqlite"]).write_record_sqlite
    real_excel = __import__("src.writer.api", fromlist=["write_record"]).write_record

    def track_sqlite(record, *args, **kwargs):
        order.append(f"sqlite:{record.assay_key}")
        return real_sqlite(record, *args, **kwargs)

    def track_excel(record, ruleset, output_dir, **kwargs):
        order.append(f"excel:{record.assay_key}")
        if record.assay_key == "(2222)" and order.count("excel:(2222)") == 1:
            raise RuntimeError("missing sheet")
        return real_excel(record, ruleset, output_dir, **kwargs)

    monkeypatch.setattr("src.dbwriter.api.write_record_sqlite", track_sqlite)
    monkeypatch.setattr("src.writer.api.write_record", track_excel)
    first = submit(str(pdf), str(root), output_mode="both", sqlite_path=str(sqlite_path), device_id="dev1")
    assert first.status == "DONE"
    assert order[:2] == ["sqlite:(1111)", "sqlite:(2222)"]
    assert _run_count(sqlite_path) == 2
    ids = _run_ids(sqlite_path)
    release_a = release_validated_run(
        root, first.job_id, ids[0], operator_initials="AB",
        validated_at_utc="2026-09-25T10:00:00+00:00", validation_id=1,
    )
    release_b = release_validated_run(
        root, first.job_id, ids[1], operator_initials="CD",
        validated_at_utc="2026-09-25T10:01:00+00:00", validation_id=2,
    )
    assert release_a.status == "DONE"
    assert release_b.status == "FAILED"
    before = order[:]
    second = retry_excel_export(root, first.job_id)
    assert second.status == "DONE"
    assert order[len(before) :] == ["excel:(2222)"]
    assert _run_count(sqlite_path) == 2
    plan = _state(root, first.job_id)["output_plan"]["assays"]
    assert {item["excel_status"] for item in plan} == {"success"}


def test_legacy_failed_without_plan_uses_full_path_and_done_stays_skipped(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    root = tmp_path / "proj"
    _write_project(root, assays=[("(1111)", "Assay A", "AssayA.json", "{assay_name}.xlsx")])
    pdf = root / "legacy.pdf"
    pdf.write_bytes(b"legacy-failed")
    monkeypatch.setattr("src.parser.api.parse", lambda path: _parse_one(Path(path)))
    parsed = {"count": 0}
    original = __import__("src.parser.api", fromlist=["parse"]).parse

    def counting_parse(path):
        parsed["count"] += 1
        return original(path)

    monkeypatch.setattr("src.parser.api.parse", counting_parse)
    first = submit(str(pdf), str(root), output_mode="sqlite", sqlite_path=str(root / "output" / "final" / "res.sqlite3"))
    assert first.status == "DONE"
    done_again = submit(str(pdf), str(root), output_mode="sqlite", sqlite_path=str(root / "output" / "final" / "res.sqlite3"))
    assert done_again.status == "SKIPPED"
    assert done_again.details["reason"] == "already_done"

    legacy_pdf = root / "legacy-failed.pdf"
    legacy_pdf.write_bytes(b"legacy-failed-body")
    monkeypatch.setattr("src.parser.api.parse", lambda path: _parse_one(Path(path)))
    preview = submit(
        str(legacy_pdf),
        str(root),
        output_mode="sqlite",
        sqlite_path=str(root / "output" / "final" / "legacy.sqlite3"),
    )
    state_path = root / "jobs" / f"{preview.job_id}.json"
    state = json.loads(state_path.read_text(encoding="utf-8"))
    state["status"] = "FAILED"
    state["error"] = "excel_write_failed:(1111):old"
    state["partial_writes"] = state.get("steps", [{}])[-1].get("writes", [])
    state.pop("output_plan", None)
    state.pop("export_retry_available", None)
    state_path.write_text(json.dumps(state), encoding="utf-8")
    parsed["count"] = 0
    monkeypatch.setattr("src.parser.api.parse", counting_parse)
    retried = submit(
        str(legacy_pdf),
        str(root),
        output_mode="sqlite",
        sqlite_path=str(root / "output" / "final" / "legacy.sqlite3"),
    )
    assert retried.status == "DONE"
    assert parsed["count"] == 1


def test_output_modes_and_duplicate_pending_stay_compatible(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    root = tmp_path / "proj"
    _write_project(
        root,
        assays=[
            ("(1111)", "Assay A", "AssayA.json", "A.xlsx"),
            ("(2222)", "Assay B", "AssayB.json", "B.xlsx"),
        ],
    )
    sqlite_path = root / "output" / "final" / "res.sqlite3"
    monkeypatch.setattr("src.parser.api.parse", lambda path: _parse_one(Path(path)))
    pdf_sqlite = root / "sqlite.pdf"
    pdf_sqlite.write_bytes(b"sqlite-only")
    sqlite_result = submit(str(pdf_sqlite), str(root), output_mode="sqlite", sqlite_path=str(sqlite_path), device_id="dev1")
    assert sqlite_result.status == "DONE"
    sqlite_state = _state(root, sqlite_result.job_id)
    assert sqlite_state["sqlite_status"] == "success"
    assert sqlite_state["structured_status"] == "success"
    assert sqlite_state["excel_status"] == "not_required"
    assert sqlite_state["export_retry_available"] is False
    assert "output_plan" not in sqlite_state
    assert not (root / "output" / "final" / "A.xlsx").exists()

    monkeypatch.setattr("src.writer.api.write_record", lambda *_a, **_k: (_ for _ in ()).throw(RuntimeError("nope")))
    pdf_excel = root / "excel.pdf"
    pdf_excel.write_bytes(b"excel-only")
    excel_result = submit(str(pdf_excel), str(root), output_mode="excel", device_id="dev1")
    assert excel_result.status == "FAILED"
    assert str(excel_result.details["error"]).startswith("excel_write_failed:")
    assert not str(excel_result.details["error"]).startswith("excel_write_failed_after_save:")
    excel_state = _state(root, excel_result.job_id)
    assert excel_state["output_plan"]["version"] == 2
    assert excel_state["output_plan"]["output_mode"] == "excel"
    assert excel_state["output_plan"]["sqlite_status"] == "not_required"
    assert excel_state["sqlite_status"] == "not_required"
    assert excel_state["structured_status"] == "not_required"
    assert excel_state["export_retry_available"] is True
    from src.application.api import friendly_processing_message

    friendly = friendly_processing_message(str(excel_result.details["error"]))
    assert friendly == "Excel-Ausgabe fehlgeschlagen"
    assert "gespeichert" not in friendly.lower()

    monkeypatch.undo()
    monkeypatch.setattr("src.parser.api.parse", lambda path: _parse_two(Path(path)))
    pdf_first = root / "dup1.pdf"
    pdf_second = root / "dup2.pdf"
    pdf_first.write_bytes(b"dup-first")
    pdf_second.write_bytes(b"dup-second")
    first = submit(str(pdf_first), str(root), output_mode="both", sqlite_path=str(sqlite_path), device_id="dev1")
    assert first.status == "DONE"

    def parse_duplicate(path):
        text = Path(path).read_bytes()
        if text == b"dup-mixed":
            lines = _lines("Assay A", "(1111)", "LOTA", "PLATE1", "TESTX") + _lines(
                "Assay B", "(2222)", "LOTNEW", "PLATENEW", "TESTNEW"
            )
            return ParsedDocument(source_path=str(path), pages=[ParsedPage(page_number=1, lines=lines)], meta={"page_count": 1})
        return _parse_two(Path(path))

    monkeypatch.setattr("src.parser.api.parse", parse_duplicate)
    pdf_mixed = root / "mixed.pdf"
    pdf_mixed.write_bytes(b"dup-mixed")
    mixed = submit(str(pdf_mixed), str(root), output_mode="both", sqlite_path=str(sqlite_path), device_id="dev1")
    assert mixed.status == "DONE"
    by_key = {item["assay_key"]: item for item in mixed.details["writes"]}
    assert by_key["(1111)"]["duplicate_status"] == "pending"
    assert any(out["status"] == "skipped_duplicate_pending" for out in by_key["(1111)"]["outputs"])
    excel_outputs = [out for out in by_key["(2222)"]["outputs"] if out.get("sink") == "excel"]
    assert excel_outputs == []
    mixed_plan = _state(root, mixed.job_id)["output_plan"]["assays"]
    statuses = {item["assay_record"]["assay_key"]: item["excel_status"] for item in mixed_plan}
    assert statuses["(1111)"] == "skipped_duplicate_pending"
    assert statuses["(2222)"] == "awaiting_validation"


def test_sqlite_failure_does_not_enable_export_only_retry(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    root = tmp_path / "proj"
    _write_project(root, assays=[("(1111)", "Assay A", "AssayA.json", "{assay_name}.xlsx")])
    pdf = root / "sqlite-fail.pdf"
    pdf.write_bytes(b"sqlite-fail")
    calls = {"parse": 0, "sqlite": 0, "excel": 0}
    monkeypatch.setattr(
        "src.parser.api.parse",
        lambda path: calls.__setitem__("parse", calls["parse"] + 1) or _parse_one(Path(path)),
    )
    monkeypatch.setattr(
        "src.dbwriter.api.write_record_sqlite",
        lambda *_a, **_k: calls.__setitem__("sqlite", calls["sqlite"] + 1) or (_ for _ in ()).throw(RuntimeError("locked")),
    )
    monkeypatch.setattr(
        "src.writer.api.write_record",
        lambda *_a, **_k: calls.__setitem__("excel", calls["excel"] + 1) or (_ for _ in ()).throw(AssertionError("excel")),
    )
    first = submit(str(pdf), str(root), output_mode="both", sqlite_path=str(root / "output" / "final" / "res.sqlite3"))
    assert first.status == "FAILED"
    assert str(first.details["error"]).startswith("sqlite_write_failed:")
    state = _state(root, first.job_id)
    assert state["sqlite_status"] == "failed"
    assert state.get("export_retry_available") is not True
    assert state["output_plan"]["version"] == 2
    assert state["output_plan"]["assays"][0]["sqlite_status"] == "failed"
    second = submit(
        str(pdf),
        str(root),
        output_mode="both",
        sqlite_path=str(root / "output" / "final" / "other.sqlite3"),
        device_id="other",
    )
    assert second.status == "FAILED"
    assert str(second.details["error"]).startswith("sqlite_write_failed:")
    assert calls["parse"] == 2
    assert calls["sqlite"] == 2
    assert calls["excel"] == 0


def test_multi_assay_validation_releases_only_selected_run(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    root = tmp_path / "proj"
    _write_project(
        root,
        assays=[
            ("(1111)", "Assay A", "AssayA.json", "A.xlsx"),
            ("(2222)", "Assay B", "AssayB.json", "B.xlsx"),
        ],
    )
    pdf = root / "crash.pdf"
    pdf.write_bytes(b"crash-mid-excel")
    sqlite_path = root / "output" / "final" / "res.sqlite3"
    excel_calls: list[str] = []
    monkeypatch.setattr("src.parser.api.parse", lambda path: _parse_two(Path(path)))
    real_excel = __import__("src.writer.api", fromlist=["write_record"]).write_record

    def track_excel(record, ruleset, output_dir, **kwargs):
        excel_calls.append(record.assay_key)
        return real_excel(record, ruleset, output_dir, **kwargs)

    monkeypatch.setattr("src.writer.api.write_record", track_excel)
    submitted = submit(str(pdf), str(root), output_mode="both", sqlite_path=str(sqlite_path), device_id="dev1")
    assert submitted.status == "DONE"

    state = json.loads(next((root / "jobs").glob("*.json")).read_text(encoding="utf-8"))
    statuses = {
        item["assay_record"]["assay_key"]: item["excel_status"]
        for item in state["output_plan"]["assays"]
    }
    assert statuses["(1111)"] == "awaiting_validation"
    assert statuses["(2222)"] == "awaiting_validation"
    assert state["sqlite_status"] == "success"
    assert state["export_retry_available"] is False
    assert not str(state.get("error") or "").startswith("excel_write_failed")
    assert _run_count(sqlite_path) == 2
    ids = _run_ids(sqlite_path)
    assert excel_calls == []
    released = release_validated_run(
        root, submitted.job_id, ids[0], operator_initials="AB",
        validated_at_utc="2026-09-25T10:00:00+00:00", validation_id=1,
    )
    assert released.status == "DONE"
    assert excel_calls == ["(1111)"]
    assert _run_ids(sqlite_path) == ids
    done = _state(root, submitted.job_id)
    assert done["output_plan"]["output_mode"] == "both"
    assert done["sqlite_status"] == "success"
    final_statuses = {
        item["assay_record"]["assay_key"]: item["excel_status"]
        for item in done["output_plan"]["assays"]
    }
    assert final_statuses == {"(1111)": "success", "(2222)": "awaiting_validation"}


def test_awaiting_validation_plan_is_terminal_before_excel(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    root = tmp_path / "proj"
    _write_project(root, assays=[("(1111)", "Assay A", "AssayA.json", "Frozen.xlsx")])
    pdf = root / "terminal.pdf"
    pdf.write_bytes(b"terminal-excel")
    sqlite_path = root / "output" / "final" / "res.sqlite3"
    monkeypatch.setattr("src.parser.api.parse", lambda path: _parse_one(Path(path)))
    submitted = submit(str(pdf), str(root), output_mode="both", sqlite_path=str(sqlite_path), device_id="dev1")
    assert submitted.status == "DONE"
    state = json.loads(next((root / "jobs").glob("*.json")).read_text(encoding="utf-8"))
    assert state["status"] == "DONE"
    assert state["excel_status"] == "awaiting_validation"
    assert state["export_retry_available"] is False
    assert state.get("error", "") == ""
    ids = _run_ids(sqlite_path)
    assert len(ids) == 1

    def forbidden(name):
        def _forbidden(*_args, **_kwargs):
            raise AssertionError(name)

        return _forbidden

    monkeypatch.setattr("src.parser.api.parse", forbidden("parse"))
    monkeypatch.setattr("src.normalizer.api.normalize_lines", forbidden("normalize"))
    monkeypatch.setattr("src.assaychooser.api.detect_assays", forbidden("chooser"))
    monkeypatch.setattr("src.contentsplitter.api.split_by_assay_name_and_key", forbidden("splitter"))
    monkeypatch.setattr("src.extractor.api.extract_record", forbidden("extract"))
    monkeypatch.setattr("src.dbwriter.api.write_record_sqlite", forbidden("sqlite"))
    monkeypatch.setattr("src.writer.api.write_record", forbidden("writer"))
    follow = submit(str(pdf), str(root), output_mode="both", sqlite_path=str(sqlite_path), device_id="dev1")
    assert follow.status == "SKIPPED"
    assert follow.details["reason"] == "already_done"
    assert _run_ids(sqlite_path) == ids


def test_crash_after_excel_save_is_retryable_and_dedupe_idempotent(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    root = tmp_path / "proj"
    _write_project(root, assays=[("(1111)", "Assay A", "AssayA.json", "Crash.xlsx")])
    pdf = root / "crash-after-save.pdf"
    pdf.write_bytes(b"crash-after-save")
    sqlite_path = root / "output" / "final" / "res.sqlite3"
    monkeypatch.setattr("src.parser.api.parse", lambda path: _parse_one(Path(path)))
    submitted = submit(str(pdf), str(root), output_mode="both", sqlite_path=str(sqlite_path), device_id="dev1")
    run_id = _run_ids(sqlite_path)[0]
    real_excel = __import__("src.writer.api", fromlist=["write_record"]).write_record

    def save_then_abort(record, ruleset, output_dir, **kwargs):
        real_excel(record, ruleset, output_dir, **kwargs)
        raise SystemExit(9)

    monkeypatch.setattr("src.writer.api.write_record", save_then_abort)
    with pytest.raises(SystemExit):
        release_validated_run(
            root, submitted.job_id, run_id, operator_initials="AB",
            validated_at_utc="2026-09-25T10:00:00+00:00", validation_id=1,
        )
    crashed = _state(root, submitted.job_id)
    assert crashed["status"] == "DONE"
    assert crashed["excel_status"] == "failed"
    assert crashed["export_retry_available"] is True
    assert crashed["output_plan"]["assays"][0]["validation"]["operator_initials"] == "AB"

    pdf.unlink()
    monkeypatch.setattr("src.writer.api.write_record", real_excel)
    monkeypatch.setattr("src.parser.api.parse", lambda *_a, **_k: (_ for _ in ()).throw(AssertionError("parse")))
    monkeypatch.setattr(
        "src.dbwriter.api.write_record_sqlite",
        lambda *_a, **_k: (_ for _ in ()).throw(AssertionError("sqlite")),
    )
    retried = retry_excel_export(root, submitted.job_id)
    assert retried.status == "DONE"
    wb = load_workbook(root / "output" / "final" / "Crash.xlsx")
    ws = wb["LOTA"]
    headers = [cell.value for cell in ws[1]]
    assert ws.max_row == 2
    assert ws.cell(2, headers.index("VALIDIERT_DURCH") + 1).value == "AB"
    assert ws.cell(2, headers.index("VALIDIERT_AM") + 1).value == "2026-09-25T10:00:00+00:00"


def test_multi_assay_sqlite_commit_before_state_save_resumes_without_duplicate(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    root = tmp_path / "proj"
    _write_project(
        root,
        assays=[
            ("(1111)", "Assay A", "AssayA.json", "A.xlsx"),
            ("(2222)", "Assay B", "AssayB.json", "B.xlsx"),
        ],
    )
    pdf = root / "sqlite-crash.pdf"
    pdf.write_bytes(b"sqlite-crash")
    sqlite_path = root / "output" / "final" / "res.sqlite3"
    monkeypatch.setattr("src.parser.api.parse", lambda path: _parse_two(Path(path)))
    original_save = JobController._save_state
    crashed = {"done": False}

    def abort_after_first_commit(self, path, state):
        plan = state.get("output_plan") if isinstance(state, dict) else None
        assays = plan.get("assays") if isinstance(plan, dict) else None
        if (
            not crashed["done"]
            and sqlite_path.exists()
            and _run_count(sqlite_path) == 1
            and isinstance(assays, list)
            and assays[0].get("sqlite_status") == "success"
        ):
            crashed["done"] = True
            raise SystemExit(17)
        return original_save(self, path, state)

    monkeypatch.setattr(JobController, "_save_state", abort_after_first_commit)
    with pytest.raises(SystemExit):
        submit(str(pdf), str(root), output_mode="both", sqlite_path=str(sqlite_path), device_id="dev1")
    persisted = _state(root, next((root / "jobs").glob("*.json")).stem)
    assert persisted["output_plan"]["assays"][0]["sqlite_status"] == "pending"
    assert persisted["output_plan"]["assays"][0]["sqlite_run_id"] is None
    assert _run_count(sqlite_path) == 1

    monkeypatch.setattr(JobController, "_save_state", original_save)
    resumed = submit(str(pdf), str(root), output_mode="both", sqlite_path=str(sqlite_path), device_id="dev1")
    assert resumed.status == "DONE"
    state = _state(root, resumed.job_id)
    entries = state["output_plan"]["assays"]
    assert [entry["sqlite_status"] for entry in entries] == ["success", "success"]
    assert all(entry["sqlite_run_id"] for entry in entries)
    assert entries[0]["write_item"]["outputs"][0]["status"] == "already_persisted"
    assert _run_count(sqlite_path) == 2
    with sqlite3.connect(sqlite_path) as conn:
        assert conn.execute("SELECT COUNT(*) FROM duplicate_candidates").fetchone()[0] == 0


def test_validation_commit_then_export_lock_failure_is_visible_and_recoverable(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    from src.jobcontroller.api import get_job_evidence
    from src.application.api import create_desktop_services
    from src.jobqueue.api import claim_next_job, enqueue_pdf_job, mark_job_done
    from src.resultvalidation.api import validate_run
    from src.runtime.api import release_exclusive, try_acquire_exclusive

    root = tmp_path / "proj"
    _write_project(root, assays=[("(1111)", "Assay A", "AssayA.json", "Locked.xlsx")])
    pdf = root / "locked.pdf"
    pdf.write_bytes(b"locked")
    queued = enqueue_pdf_job(str(root), str(pdf), source="test-app-manual")
    assert claim_next_job(str(root), "w-test") is not None
    sqlite_path = root / "output" / "final" / "res.sqlite3"
    monkeypatch.setattr("src.parser.api.parse", lambda path: _parse_one(Path(path)))
    submitted = submit(str(pdf), str(root), output_mode="both", sqlite_path=str(sqlite_path), device_id="dev1")
    assert submitted.job_id == queued.job_id
    mark_job_done(str(root), queued.job_id, worker_id="w-test")
    run_id = _run_ids(sqlite_path)[0]
    validation = validate_run({"driver": "sqlite", "path": str(sqlite_path)}, run_id, "AB")
    export_lock = root / "locks" / f"{submitted.job_id}.excel-export.lock"
    assert try_acquire_exclusive(export_lock, 900.0)
    try:
        blocked = release_validated_run(
            root,
            submitted.job_id,
            run_id,
            operator_initials=validation.operator_initials,
            validated_at_utc=validation.validated_at_utc,
            validation_id=validation.validation_id,
        )
        assert blocked.status == "FAILED"
        assert blocked.details["error"] == "excel_export_locked"
        evidence = get_job_evidence(root, submitted.job_id)
        assert evidence is not None
        assert evidence.retry_kind == "export_only"
        assert evidence.error == "excel_export_reconcile_required"
        diagnoses = create_desktop_services(root).extraction.list_job_diagnoses()
        assert len(diagnoses) == 1
        assert diagnoses[0].retry_kind == "export_only"
        state = _state(root, submitted.job_id)
        assert "validation" not in state["output_plan"]["assays"][0]
    finally:
        release_exclusive(export_lock)

    pdf.unlink()
    recovered = retry_excel_export(root, submitted.job_id)
    assert recovered.status == "DONE"
    state = _state(root, submitted.job_id)
    entry = state["output_plan"]["assays"][0]
    assert entry["validation"]["validation_id"] == validation.validation_id
    assert entry["excel_status"] == "success"
    assert create_desktop_services(root).extraction.list_job_diagnoses() == []


def test_effective_validation_reconcile_exports_without_prior_handoff(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    from src.resultvalidation.api import validate_run

    root = tmp_path / "proj"
    _write_project(root, assays=[("(1111)", "Assay A", "AssayA.json", "Reconcile.xlsx")])
    pdf = root / "reconcile.pdf"
    pdf.write_bytes(b"reconcile")
    sqlite_path = root / "output" / "final" / "res.sqlite3"
    monkeypatch.setattr("src.parser.api.parse", lambda path: _parse_one(Path(path)))
    submitted = submit(str(pdf), str(root), output_mode="both", sqlite_path=str(sqlite_path), device_id="dev1")
    run_id = _run_ids(sqlite_path)[0]
    validation = validate_run({"driver": "sqlite", "path": str(sqlite_path)}, run_id, "AB")
    pdf.unlink()

    reconciled = retry_excel_export(root, submitted.job_id)
    assert reconciled.status == "DONE"
    entry = _state(root, submitted.job_id)["output_plan"]["assays"][0]
    assert entry["validation"]["validation_id"] == validation.validation_id
    assert entry["excel_status"] == "success"
    assert (root / "output" / "final" / "Reconcile.xlsx").is_file()


def test_out_of_order_validation_release_never_overwrites_newer_correction(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    from src.resultvalidation.api import correct_validation, validate_run

    root = tmp_path / "proj"
    _write_project(root, assays=[("(1111)", "Assay A", "AssayA.json", "Monotonic.xlsx")])
    pdf = root / "monotonic.pdf"
    pdf.write_bytes(b"monotonic")
    sqlite_path = root / "output" / "final" / "res.sqlite3"
    monkeypatch.setattr("src.parser.api.parse", lambda path: _parse_one(Path(path)))
    submitted = submit(str(pdf), str(root), output_mode="both", sqlite_path=str(sqlite_path), device_id="dev1")
    run_id = _run_ids(sqlite_path)[0]
    first = validate_run({"driver": "sqlite", "path": str(sqlite_path)}, run_id, "AB")
    assert release_validated_run(
        root, submitted.job_id, run_id,
        operator_initials=first.operator_initials,
        validated_at_utc=first.validated_at_utc,
        validation_id=first.validation_id,
    ).status == "DONE"
    newest = correct_validation({"driver": "sqlite", "path": str(sqlite_path)}, run_id, "CD")
    out_of_order = release_validated_run(
        root, submitted.job_id, run_id,
        operator_initials=first.operator_initials,
        validated_at_utc=first.validated_at_utc,
        validation_id=first.validation_id,
        correction=True,
    )
    assert out_of_order.status == "DONE"
    entry = _state(root, submitted.job_id)["output_plan"]["assays"][0]
    assert entry["validation"]["validation_id"] == newest.validation_id
    wb = load_workbook(root / "output" / "final" / "Monotonic.xlsx")
    ws = wb["LOTA"]
    headers = [cell.value for cell in ws[1]]
    assert ws.max_row == 2
    assert ws.cell(2, headers.index("VALIDIERT_DURCH") + 1).value == "CD"

    stale_again = release_validated_run(
        root, submitted.job_id, run_id,
        operator_initials=first.operator_initials,
        validated_at_utc=first.validated_at_utc,
        validation_id=first.validation_id,
        correction=True,
    )
    assert stale_again.details["validation_id"] == newest.validation_id
