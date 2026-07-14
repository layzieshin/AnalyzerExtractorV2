from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook

from src.extractor.model import AssayRecord
from src.ruleresolver.model import RuleSet
from src.writer.api import write_record
from src.writer.excel_table import headers_table_eligible, sync_worksheet_table
from src.writer.writer import WriterError


def _ruleset() -> RuleSet:
    return RuleSet(
        assay_key="(1111)",
        ruleset_file="A.json",
        data={
            "assay_name": "Assay A",
            "assay_key": "(1111)",
            "lot_rule": {"regex": "x"},
            "extract_rules": {"fields": []},
            "excel_rules": {
                "excel_filename_template": "{assay_name}.xlsx",
                "sheetname_template": "{lot_id}",
                "column_mapping": {},
            },
        },
    )


def test_writer_created_then_skipped_on_same_dedupe(tmp_path: Path):
    out = tmp_path / "final"
    rec = AssayRecord(
        assay_key="(1111)",
        lot_id="LOT1",
        dedupe_key="K1",
        data={"date": "2026-01-01", "time": "10:00:00"},
        device_id="dev1",
        dedupe_version="v2",
    )
    r1 = write_record(rec, _ruleset(), str(out))
    r2 = write_record(rec, _ruleset(), str(out))
    assert r1.status == "created"
    assert r2.status == "skipped"
    wb = load_workbook(r1.excel_path)
    headers = [cell.value for cell in wb[r1.sheet_name][1]]
    assert headers[:5] == ["assay_key", "device_id", "lot_id", "dedupe_key", "dedupe_version"]


def test_writer_lock_file_blocks_write(tmp_path: Path):
    out = tmp_path / "final"
    out.mkdir(parents=True)
    (out / ".excel_writer.lock").write_text("", encoding="utf-8")
    rec = AssayRecord(assay_key="(1111)", lot_id="LOT1", dedupe_key="K1", data={})
    with pytest.raises(WriterError):
        write_record(rec, _ruleset(), str(out))


def test_writer_first_write_creates_excel_table(tmp_path: Path):
    out = tmp_path / "final"
    rec = AssayRecord(
        assay_key="(1111)",
        lot_id="LOT1",
        dedupe_key="K1",
        data={"date": "2026-01-01"},
        device_id="dev1",
        dedupe_version="v2",
    )
    result = write_record(rec, _ruleset(), str(out))
    wb = load_workbook(result.excel_path)
    ws = wb[result.sheet_name]
    assert ws.tables
    table = next(iter(ws.tables.values()))
    assert table.displayName.startswith("ARE_")
    assert table.ref.endswith("2")


def test_writer_second_write_extends_table_ref(tmp_path: Path):
    out = tmp_path / "final"
    rs = _ruleset()
    rec1 = AssayRecord(
        assay_key="(1111)",
        lot_id="LOT1",
        dedupe_key="K1",
        data={"date": "2026-01-01"},
        device_id="dev1",
        dedupe_version="v2",
    )
    rec2 = AssayRecord(
        assay_key="(1111)",
        lot_id="LOT1",
        dedupe_key="K2",
        data={"date": "2026-01-02"},
        device_id="dev1",
        dedupe_version="v2",
    )
    r1 = write_record(rec1, rs, str(out))
    write_record(rec2, rs, str(out))
    wb = load_workbook(r1.excel_path)
    ws = wb[r1.sheet_name]
    table = next(iter(ws.tables.values()))
    assert table.ref.endswith("3")
    assert ws.max_row == 3


def test_writer_duplicate_skip_keeps_row_count_and_table(tmp_path: Path):
    out = tmp_path / "final"
    rec = AssayRecord(
        assay_key="(1111)",
        lot_id="LOT1",
        dedupe_key="K1",
        data={"date": "2026-01-01"},
        device_id="dev1",
        dedupe_version="v2",
    )
    rs = _ruleset()
    r1 = write_record(rec, rs, str(out))
    write_record(rec, rs, str(out))
    wb = load_workbook(r1.excel_path)
    ws = wb[r1.sheet_name]
    assert ws.max_row == 2
    assert ws.tables


def test_writer_duplicate_headers_skip_table_sync(tmp_path: Path):
    wb = Workbook()
    ws = wb.active
    ws.title = "LOT1"
    ws.append(["assay_key", "assay_key", "lot_id"])
    ws.append(["(1111)", "dup", "LOT1"])
    sync_worksheet_table(ws, wb, ["assay_key", "assay_key", "lot_id"], "LOT1")
    assert not ws.tables


def test_headers_table_eligible_rejects_empty_and_duplicate():
    assert headers_table_eligible(["a", "b"]) is True
    assert headers_table_eligible(["a", "a"]) is False
    assert headers_table_eligible(["", "b"]) is False


def test_writer_table_name_fallback_are_table():
    wb = Workbook()
    ws = wb.active
    ws.append(["assay_key", "lot_id"])
    ws.append(["(1111)", "LOT1"])
    sync_worksheet_table(ws, wb, ["assay_key", "lot_id"], "---")
    table = next(iter(ws.tables.values()))
    assert table.displayName == "ARE_Table"


def test_writer_header_only_sheet_has_no_table(tmp_path: Path):
    wb = Workbook()
    ws = wb.active
    ws.title = "LOT1"
    ws.append(["assay_key", "lot_id"])
    sync_worksheet_table(ws, wb, ["assay_key", "lot_id"], "LOT1")
    assert not ws.tables


def test_writer_workbook_unique_are_table_names(tmp_path: Path):
    out = tmp_path / "final"
    rs = _ruleset()
    rec1 = AssayRecord(
        assay_key="(1111)",
        lot_id="LOT1",
        dedupe_key="K1",
        data={"date": "2026-01-01"},
        device_id="dev1",
        dedupe_version="v2",
    )
    rec2 = AssayRecord(
        assay_key="(1111)",
        lot_id="LOT2",
        dedupe_key="K2",
        data={"date": "2026-01-02"},
        device_id="dev1",
        dedupe_version="v2",
    )
    r1 = write_record(rec1, rs, str(out))
    write_record(rec2, rs, str(out))
    wb = load_workbook(r1.excel_path)
    names = {table.displayName for ws in wb.worksheets for table in ws.tables.values()}
    assert len(names) == 2
    assert all(name.startswith("ARE_") for name in names)
